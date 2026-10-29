"""Exportadores multi-formato para relatórios do Grafana MCP Server (PDF, Excel, CSV, JSON, Markdown)."""

import csv
import json
import os
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def export_to_json(data: dict[str, Any], filepath: str) -> str:
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return filepath


def export_to_csv(headers: list[str], rows: list[list[Any]], filepath: str) -> str:
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    return filepath


def export_to_markdown(title: str, headers: list[str], rows: list[list[Any]], filepath: str) -> str:
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    lines = [f"# {title}\n", "| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(str(cell) for cell in row) + " |")
    content = "\n".join(lines)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)
    return filepath


def export_to_xlsx(title: str, sheets_data: dict[str, dict[str, Any]], filepath: str) -> str:
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb = openpyxl.Workbook()
    # Remove sheet padrão
    wb.remove(wb.active)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for sheet_name, content in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        headers = content.get("headers", [])
        rows = content.get("rows", [])

        # Adiciona título
        ws.append([f"{title} - {sheet_name}"])
        ws.cell(row=1, column=1).font = Font(name="Segoe UI", size=14, bold=True, color="1F497D")
        ws.append([])  # Linha vazia

        # Adiciona cabeçalhos
        if headers:
            ws.append(headers)
            header_row_idx = 3
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=header_row_idx, column=col_idx)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center", vertical="center")

        # Adiciona dados
        for r_idx, row in enumerate(rows, start=4):
            ws.append(row)
            for c_idx in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = border
                cell.font = Font(name="Segoe UI", size=10)

        # Auto-ajuste de largura de colunas
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(filepath)
    return filepath


def export_to_pdf(
    title: str,
    summary_text: str,
    headers: list[str],
    rows: list[list[Any]],
    filepath: str,
) -> str:
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        name="DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        textColor=colors.HexColor("#1F497D"),
        spaceAfter=12,
    )
    body_style = ParagraphStyle(
        name="DocBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        textColor=colors.HexColor("#333333"),
        spaceAfter=14,
    )

    story = [
        Paragraph(title, title_style),
        Paragraph(summary_text, body_style),
        Spacer(1, 10),
    ]

    if headers and rows:
        table_data = [headers] + [[str(c) for c in r] for r in rows[:100]]
        t = Table(table_data, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F497D")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                    ("TOPPADDING", (0, 0), (-1, 0), 6),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 1), (-1, -1), 8),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FC")]),
                ]
            )
        )
        story.append(t)

    doc.build(story)
    return filepath
